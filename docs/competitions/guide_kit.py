"""Shared DOCX toolkit for the Competition domain guides.

Every guide is built on top of the same branded Word letterhead the
platform's own generated letters use — `backend/app/static/templates/docx/
payment_letter.docx` carries the SPACE. banner in its default header part
and the copyright / spacepoint.ae / "Page X of Y" line in its default
footer part. We open that file, throw away its body (keeping the section
properties, so the header/footer references and page setup survive), and
write the guide into the empty shell. That way the guides can never drift
from the letterhead the rest of the system prints.

Screenshots are deliberately NOT embedded: each one is a labelled, shaded
placeholder box naming what to capture and what to annotate, so whoever
takes the screenshots can drop the image straight into the box.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

REPO_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = REPO_ROOT / "backend/app/static/templates/docx/payment_letter.docx"

# Brand palette — the banner colour is lifted from letterhead.py's
# BANNER_COLOR so headings sit in the same family as the header art.
BRAND = RGBColor(0x23, 0x11, 0x34)
BRAND_HEX = "231134"
ACCENT = RGBColor(0x5B, 0x3A, 0x86)
ACCENT_HEX = "5B3A86"
BODY = RGBColor(0x2B, 0x2B, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x72)
LINK = RGBColor(0x05, 0x63, 0xC1)

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

SHADE_SOFT = "F4F2F7"     # callouts / table headers
SHADE_SHOT = "FAF9FC"     # screenshot placeholder fill
SHADE_ZEBRA = "FBFAFD"    # alternating table rows


# ── low-level XML helpers ────────────────────────────────────────────────
def _shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _cell_shade(cell, fill: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), fill)


def _cell_borders(cell, *, color: str = "D9D4E3", size: int = 6,
                  style: str = "single", edges=("top", "left", "bottom", "right")) -> None:
    borders = OxmlElement("w:tcBorders")
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    cell._tc.get_or_add_tcPr().append(borders)


def _cell_margins(cell, *, top=100, bottom=100, left=140, right=140) -> None:
    mar = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    cell._tc.get_or_add_tcPr().append(mar)


def _fixed_layout(table) -> None:
    """Word only honours explicit column widths under a fixed layout."""
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    tr_pr.append(header)


def _cant_split(row) -> None:
    """Keep a one-row box (callout, screenshot placeholder) whole rather than
    letting Word break it across a page."""
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _no_table_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


# ── schema ordering ──────────────────────────────────────────────────────
# WordprocessingML property elements are a *sequence*, not a bag: Word and
# LibreOffice both refuse to open a file whose <w:pPr>/<w:tcPr>/<w:tblPr>
# children are out of order. python-docx only enforces that for the
# properties it sets itself, so everything we append by hand above gets
# sorted back into schema order in one pass at save time.
_PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]
_TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "cellIns", "cellDel",
    "cellMerge", "tcPrChange",
]
_TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
    "tblPrChange",
]
_ORDERS = {"pPr": _PPR_ORDER, "tcPr": _TCPR_ORDER, "tblPr": _TBLPR_ORDER}


def normalize_order(root) -> None:
    for name, order in _ORDERS.items():
        for pr in root.iter(qn(f"w:{name}")):
            children = list(pr)
            keyed = sorted(
                enumerate(children),
                key=lambda pair: (
                    order.index(pair[1].tag.split("}")[-1])
                    if pair[1].tag.split("}")[-1] in order else len(order),
                    pair[0],
                ),
            )
            for _, child in keyed:
                pr.append(child)


# ── document shell ───────────────────────────────────────────────────────
class Guide:
    """One guide document. Screenshot numbering is per-document."""

    def __init__(self, title: str, subtitle: str, meta: list[tuple[str, str]]):
        self.doc = Document(str(TEMPLATE))
        self._clear_body()
        self._set_defaults()
        self.shot_no = 0
        self._title_block(title, subtitle, meta)

    def _clear_body(self) -> None:
        body = self.doc.element.body
        for child in list(body):
            if child.tag == qn("w:sectPr"):
                continue
            body.remove(child)

    def _set_defaults(self) -> None:
        # This template names its default style "normal" (lower-case), so look
        # it up by that name and fall back to the usual spelling.
        try:
            normal = self.doc.styles["normal"]
        except KeyError:
            normal = self.doc.styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = BODY
        pf = normal.paragraph_format
        pf.space_after = Pt(7)
        pf.line_spacing = 1.15

    # ── text primitives ──────────────────────────────────────────────
    def _runs(self, paragraph, text: str, *, size=None, color=None, bold=False, italic=False):
        """Tiny inline markup: **bold**, *italic* and `code`."""
        for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`)", text):
            if not chunk:
                continue
            run = paragraph.add_run()
            if chunk.startswith("**") and chunk.endswith("**"):
                run.text = chunk[2:-2]
                run.bold = True
            elif chunk.startswith("*") and chunk.endswith("*"):
                run.text = chunk[1:-1]
                run.italic = True
            elif chunk.startswith("`") and chunk.endswith("`"):
                run.text = chunk[1:-1]
                run.font.name = MONO_FONT
                run.font.size = Pt((size or 10.5) - 0.5)
                run.font.color.rgb = ACCENT
            else:
                run.text = chunk
                run.bold = bold
            run.italic = run.italic or italic
            if run.font.size is None and size:
                run.font.size = Pt(size)
            if run.font.color.rgb is None and color is not None:
                run.font.color.rgb = color
            if run.font.name is None:
                run.font.name = BODY_FONT
        return paragraph

    def para(self, text: str = "", *, size=10.5, color=BODY, italic=False,
             space_after=7, space_before=0, indent=0.0, align=None):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if indent:
            p.paragraph_format.left_indent = Inches(indent)
        if align is not None:
            p.alignment = align
        if text:
            self._runs(p, text, size=size, color=color, italic=italic)
        return p

    def title(self, text: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = BODY_FONT
        run.font.size = Pt(26)
        run.bold = True
        run.font.color.rgb = BRAND
        return p

    def h1(self, text: str, *, page_break: bool = True):
        if page_break:
            self.page_break()
        p = self.doc.add_paragraph(style="Heading 1")
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.font.name = BODY_FONT
        run.font.size = Pt(17)
        run.bold = True
        run.font.color.rgb = BRAND
        self._rule(p)
        _keep_with_next(p)
        return p

    def h2(self, text: str):
        p = self.doc.add_paragraph(style="Heading 2")
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(text)
        run.font.name = BODY_FONT
        run.font.size = Pt(13)
        run.bold = True
        run.font.color.rgb = ACCENT
        _keep_with_next(p)
        return p

    def h3(self, text: str):
        p = self.doc.add_paragraph(style="Heading 3")
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = BODY_FONT
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = BODY
        _keep_with_next(p)
        return p

    def _rule(self, paragraph, color: str = "D9D4E3") -> None:
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), color)
        pbdr.append(bottom)
        paragraph._p.get_or_add_pPr().append(pbdr)

    def page_break(self) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    # ── lists ────────────────────────────────────────────────────────
    def bullets(self, items: list[str], *, indent=0.22, size=10.5):
        for item in items:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(indent + 0.18)
            pf.first_line_indent = Inches(-0.18)
            pf.space_after = Pt(3)
            dot = p.add_run("•\t")
            dot.font.color.rgb = ACCENT
            dot.bold = True
            dot.font.size = Pt(size)
            self._runs(p, item, size=size)
        self.para(space_after=2)

    def steps(self, items: list[str], *, start: int = 1, size=10.5):
        """Explicitly numbered steps — a printed manual reads better with the
        number visible than with Word's list numbering, and it survives
        copy/paste into a ticket."""
        for i, item in enumerate(items, start=start):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.36)
            pf.first_line_indent = Inches(-0.36)
            pf.space_after = Pt(5)
            num = p.add_run(f"{i}.\t")
            num.bold = True
            num.font.color.rgb = ACCENT
            num.font.size = Pt(size)
            self._runs(p, item, size=size)
        self.para(space_after=2)

    # ── boxes ────────────────────────────────────────────────────────
    def callout(self, kind: str, text: str):
        """kind: Note / Watch out / Tip / Rule."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        cell = table.cell(0, 0)
        _cell_shade(cell, SHADE_SOFT)
        _cell_margins(cell, top=120, bottom=120)
        _cell_borders(cell, color=ACCENT_HEX, size=18, edges=("left",))
        _cell_borders(cell, color="E4DFEC", size=4, edges=("top", "bottom", "right"))
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        label = p.add_run(f"{kind.upper()}   ")
        label.bold = True
        label.font.size = Pt(8.5)
        label.font.color.rgb = ACCENT
        label.font.name = BODY_FONT
        self._runs(p, text, size=10)
        _cant_split(table.rows[0])
        self.para(space_after=4)
        return table

    def screenshot(self, caption: str, *, where: str = "", annotate: list[str] | None = None):
        self.shot_no += 1
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        cell = table.cell(0, 0)
        _cell_shade(cell, SHADE_SHOT)
        _cell_margins(cell, top=160, bottom=160, left=180, right=180)
        _cell_borders(cell, color=ACCENT_HEX, size=6, style="dashed")

        head = cell.paragraphs[0]
        head.paragraph_format.space_after = Pt(2)
        tag = head.add_run(f"SCREENSHOT {self.shot_no}")
        tag.bold = True
        tag.font.size = Pt(9)
        tag.font.color.rgb = ACCENT
        tag.font.name = BODY_FONT
        rest = head.add_run(f"   {caption}")
        rest.font.size = Pt(10)
        rest.bold = True
        rest.font.color.rgb = BODY
        rest.font.name = BODY_FONT

        if where:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            self._runs(p, f"**Where:** {where}", size=9, color=MUTED)
            for run in p.runs:
                if not run.bold and run.font.color.rgb is None:
                    run.font.color.rgb = MUTED
        for note in annotate or []:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            arrow = p.add_run("↳\t")
            arrow.font.size = Pt(9)
            arrow.font.color.rgb = ACCENT
            self._runs(p, note, size=9, color=MUTED)

        drop = cell.add_paragraph()
        drop.alignment = WD_ALIGN_PARAGRAPH.CENTER
        drop.paragraph_format.space_before = Pt(14)
        drop.paragraph_format.space_after = Pt(14)
        hint = drop.add_run("[ paste the annotated screenshot here ]")
        hint.italic = True
        hint.font.size = Pt(9)
        hint.font.color.rgb = MUTED
        hint.font.name = BODY_FONT

        _cant_split(table.rows[0])
        self.para(space_after=6)
        return table

    def table(self, headers: list[str], rows: list[list[str]], *, widths: list[float] | None = None,
              size=9.5):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        _no_table_borders(table)
        table.autofit = False

        hdr = table.rows[0]
        for i, text in enumerate(headers):
            cell = hdr.cells[i]
            _cell_shade(cell, SHADE_SOFT)
            _cell_margins(cell, top=80, bottom=80)
            _cell_borders(cell, color=ACCENT_HEX, size=6, edges=("bottom",))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(size)
            run.font.color.rgb = BRAND
            run.font.name = BODY_FONT

        for r, row in enumerate(rows):
            cells = table.add_row().cells
            for i, text in enumerate(row):
                cell = cells[i]
                if r % 2 == 1:
                    _cell_shade(cell, SHADE_ZEBRA)
                _cell_margins(cell, top=70, bottom=70)
                _cell_borders(cell, color="E4DFEC", size=4, edges=("bottom",))
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                self._runs(p, text, size=size)

        if widths:
            _fixed_layout(table)
            for i, width in enumerate(widths):
                table.columns[i].width = Inches(width)
            for row in table.rows:
                for i, width in enumerate(widths):
                    row.cells[i].width = Inches(width)
        _repeat_header(table.rows[0])
        self.para(space_after=6)
        return table

    # ── front matter ─────────────────────────────────────────────────
    def _title_block(self, title: str, subtitle: str, meta: list[tuple[str, str]]) -> None:
        eyebrow = self.doc.add_paragraph()
        eyebrow.paragraph_format.space_after = Pt(1)
        run = eyebrow.add_run("SPACEPOINT PLATFORM  ·  COMPETITION DOMAIN")
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = ACCENT
        run.font.name = BODY_FONT

        self.title(title)
        sub = self.para(subtitle, size=12, color=MUTED, space_after=12)
        self._rule(sub)

        table = self.doc.add_table(rows=len(meta), cols=2)
        _no_table_borders(table)
        _fixed_layout(table)
        table.columns[0].width = Inches(1.35)
        table.columns[1].width = Inches(5.15)
        for i, (label, value) in enumerate(meta):
            key, val = table.rows[i].cells
            key.width = Inches(1.35)
            val.width = Inches(5.15)
            _cell_margins(key, top=40, bottom=40, left=0)
            _cell_margins(val, top=40, bottom=40, left=0)
            kp = key.paragraphs[0]
            kp.paragraph_format.space_after = Pt(0)
            kr = kp.add_run(label.upper())
            kr.bold = True
            kr.font.size = Pt(8)
            kr.font.color.rgb = MUTED
            kr.font.name = BODY_FONT
            vp = val.paragraphs[0]
            vp.paragraph_format.space_after = Pt(0)
            self._runs(vp, value, size=9.5)
        self.para(space_after=8)

    def save(self, path: Path) -> Path:
        normalize_order(self.doc.element.body)
        self.doc.save(str(path))
        return path
