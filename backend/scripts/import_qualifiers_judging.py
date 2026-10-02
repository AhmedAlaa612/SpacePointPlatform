"""Import Ru'ya Qualifiers detailed judging scores, timing adjustments,
criteria justifications, and team qualifications from Excel workbook into the platform.

Usage:
  python scripts/import_qualifiers_judging.py --cohort-id <COHORT_UUID> [--file <PATH>] [--dry-run] [--commit]
"""
import argparse
import asyncio
import os
import re
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

# Ensure backend root is in sys.path
backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

import openpyxl
from sqlalchemy import func, select, update
from app.db.session import AsyncSessionLocal
from app.models.competitions.judging import (
    CompetitionJudge, CompetitionRubric, CompetitionRubricCriterion, CompetitionScore
)
from app.models.sessions.cohort import Cohort
from app.models.enums import UserRole
from app.models.team import Team
from app.models.user import User

DEFAULT_EXCEL_PATH = r"C:\Users\ahmed\Downloads\RUYA_Qualifiers_detailed_judging_sheet.xlsx"

CRITERIA_SPECS = [
    ("mission_concept", "Mission Concept", 15),
    ("dashboard", "Dashboard", 15),
    ("madar_mission", "Madar Mission", 15),
    ("tinkercad_circuit", "Tinkercad Circuit", 15),
    ("tinkercad_3d_design", "Tinkercad 3D Design", 10),
    ("poster", "Poster", 10),
    ("recorded_presentation", "Recorded Presentation", 10),
    ("consistency", "Consistency", 10),
]


def clean_team_name(raw: str) -> str:
    s = re.sub(r"\s*\(\d+\)\s*$", "", raw).strip()
    return s


def strip_all_parens(s: str) -> str:
    return re.sub(r"\s*\([^)]*\)", "", s).strip()


def extract_team_id(raw: str) -> str | None:
    m = re.search(r"\((\d+)\)\s*$", raw)
    return m.group(1) if m else None


async def ensure_judge_user(db, email: str, full_name: str) -> User:
    # Canonical normalization (strip + lowercase) on both lookup and
    # insert, with a case-insensitive column compare so a legacy
    # mixed-case row is reused instead of colliding (auth compares the
    # same normalized form — storing anything else mints an account that
    # can never log in).
    norm_email = (email or "").strip().lower()
    user = await db.scalar(select(User).where(func.lower(User.email) == norm_email))
    if user is None:
        user = User(
            id=uuid.uuid4(),
            email=norm_email,
            full_name=full_name,
            password_hash="!imported_judge_no_password",
            roles=[UserRole.judge],
            status="active",
        )
        db.add(user)
        await db.flush()
    else:
        current_roles = list(user.roles or [])
        if UserRole.judge not in current_roles and "judge" not in current_roles:
            user.roles = current_roles + [UserRole.judge]
            await db.flush()
    return user


async def ensure_cohort_judge(db, cohort_id: uuid.UUID, user_id: uuid.UUID) -> CompetitionJudge:
    cj = await db.scalar(
        select(CompetitionJudge).where(
            CompetitionJudge.cohort_id == cohort_id, CompetitionJudge.user_id == user_id
        )
    )
    if cj is None:
        cj = CompetitionJudge(
            id=uuid.uuid4(),
            cohort_id=cohort_id,
            user_id=user_id,
            onboarding_sent_at=datetime.now(timezone.utc),
        )
        db.add(cj)
        await db.flush()
    return cj


async def ensure_rubric(db, program_id: uuid.UUID) -> tuple[CompetitionRubric, list[CompetitionRubricCriterion]]:
    rubric = await db.scalar(select(CompetitionRubric).where(CompetitionRubric.program_id == program_id))
    if rubric is None:
        rubric = CompetitionRubric(
            id=uuid.uuid4(),
            program_id=program_id,
            name="Ru'ya Space Race Scoring Rubric",
        )
        db.add(rubric)
        await db.flush()

    criteria = (await db.execute(
        select(CompetitionRubricCriterion)
        .where(CompetitionRubricCriterion.rubric_id == rubric.id)
        .order_by(CompetitionRubricCriterion.position)
    )).scalars().all()

    if not criteria:
        created_criteria = []
        for pos, (key, label, max_pts) in enumerate(CRITERIA_SPECS):
            c = CompetitionRubricCriterion(
                id=uuid.uuid4(),
                rubric_id=rubric.id,
                key=key,
                label=label,
                max_points=max_pts,
                position=pos,
            )
            db.add(c)
            created_criteria.append(c)
        await db.flush()
        criteria = created_criteria

    return rubric, list(criteria)


async def main():
    parser = argparse.ArgumentParser(description="Import Ru'ya Qualifiers judging scores from Excel")
    parser.add_argument("--cohort-id", required=True, help="Target cohort UUID")
    parser.add_argument("--file", default=DEFAULT_EXCEL_PATH, help="Path to Excel sheet")
    parser.add_argument("--judge1-email", default="judge1.ruya@spacepoint.ae", help="Email for Judge 1")
    parser.add_argument("--judge2-email", default="judge2.ruya@spacepoint.ae", help="Email for Judge 2")
    parser.add_argument("--commit", action="store_true", help="Write changes to database (default is dry-run)")
    parser.add_argument("--dry-run", action="store_true", help="Preview changes without writing to database (default)")
    args = parser.parse_args()

    cohort_uuid = uuid.UUID(args.cohort_id)
    excel_path = Path(args.file)
    if not excel_path.exists():
        print(f"❌ Error: Excel file not found at: {excel_path}")
        sys.exit(1)

    print(f"📂 Loading workbook: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    async with AsyncSessionLocal() as db:
        # 1. Validate cohort
        cohort = await db.get(Cohort, cohort_uuid)
        if not cohort:
            print(f"❌ Error: Cohort '{args.cohort_id}' not found in database.")
            sys.exit(1)
        print(f"✓ Cohort found: '{cohort.name}' (Program ID: {cohort.program_id})")

        # 2. Rubric
        rubric, criteria = await ensure_rubric(db, cohort.program_id)
        print(f"✓ Rubric ready: '{rubric.name}' ({len(criteria)} criteria)")

        # 3. Judges
        judge1 = await ensure_judge_user(db, args.judge1_email, "Judge 1")
        judge2 = await ensure_judge_user(db, args.judge2_email, "Judge 2")
        await ensure_cohort_judge(db, cohort_uuid, judge1.id)
        await ensure_cohort_judge(db, cohort_uuid, judge2.id)
        print(f"✓ Judge 1: {judge1.full_name} ({judge1.email})")
        print(f"✓ Judge 2: {judge2.full_name} ({judge2.email})")

        # 4. Fetch all teams in cohort
        teams = (await db.execute(
            select(Team).where(Team.cohort_id == cohort_uuid)
        )).scalars().all()
        print(f"✓ Found {len(teams)} teams in cohort.")

        # Build lookup maps for teams
        team_by_roster_id = {}
        team_by_clean_name = {}
        team_by_exact_name = {}
        team_by_base_name = {}
        for t in teams:
            team_by_exact_name[t.name.strip().lower()] = t
            team_by_clean_name[clean_team_name(t.name).lower()] = t
            team_by_base_name[strip_all_parens(t.name).lower()] = t
            if t.roster_team_id:
                team_by_roster_id[str(t.roster_team_id).strip().lower()] = t

        matched_team_ids = set()

        def find_team(raw_name: str) -> tuple[Team | None, str]:
            c_name = clean_team_name(raw_name)
            e_id = extract_team_id(raw_name)
            base_name = strip_all_parens(raw_name)

            # 1. Exact Roster ID match
            if e_id and e_id.lower() in team_by_roster_id:
                t = team_by_roster_id[e_id.lower()]
                if t.id not in matched_team_ids:
                    matched_team_ids.add(t.id)
                    return t, f"BY_ROSTER_ID ({e_id})"

            # 2. Exact full name match
            if c_name.lower() in team_by_exact_name:
                t = team_by_exact_name[c_name.lower()]
                if t.id not in matched_team_ids:
                    matched_team_ids.add(t.id)
                    return t, "BY_EXACT_NAME"

            # 3. Clean name match (excluding trailing id)
            if c_name.lower() in team_by_clean_name:
                t = team_by_clean_name[c_name.lower()]
                if t.id not in matched_team_ids:
                    matched_team_ids.add(t.id)
                    return t, "BY_CLEAN_NAME"

            # 4. Base name match (excluding all parens)
            if base_name.lower() in team_by_base_name:
                t = team_by_base_name[base_name.lower()]
                if t.id not in matched_team_ids:
                    matched_team_ids.add(t.id)
                    return t, f"BY_BASE_NAME ({base_name})"

            if raw_name.strip().lower() in team_by_exact_name:
                t = team_by_exact_name[raw_name.strip().lower()]
                if t.id not in matched_team_ids:
                    matched_team_ids.add(t.id)
                    return t, "BY_RAW_NAME"

            # 5. Safe substring matching on team NAME (min 4 chars)
            for t in teams:
                if t.id in matched_team_ids:
                    continue
                if len(c_name) >= 4 and (c_name.lower() in t.name.lower() or t.name.lower() in c_name.lower()):
                    matched_team_ids.add(t.id)
                    return t, f"BY_SUBSTRING ({t.name})"
                if len(base_name) >= 4 and (base_name.lower() in t.name.lower() or t.name.lower() in base_name.lower()):
                    matched_team_ids.add(t.id)
                    return t, f"BY_SUBSTRING ({t.name})"

            return None, "UNMATCHED"

        # 5. Parse Judge 1 Sheet
        s_j1 = wb["Judge 1 sheet"]
        j1_data_by_raw = {}
        for r in range(6, s_j1.max_row + 1):
            t_name = s_j1.cell(r, 2).value
            if not t_name:
                continue
            raw = str(t_name).strip()
            timing = float(s_j1.cell(r, 4).value or 0.0)
            eval_bonus = float(s_j1.cell(r, 5).value or 0.0)
            final_score = float(s_j1.cell(r, 6).value or 0.0)
            qual = s_j1.cell(r, 9).value
            eval_notes = s_j1.cell(r, 11).value or ""
            deductions = s_j1.cell(r, 12).value or ""
            bonus_notes = s_j1.cell(r, 13).value or ""

            notes_parts = []
            if eval_notes:
                notes_parts.append(f"Evaluation: {eval_notes}")
            if deductions:
                notes_parts.append(f"Deductions: {deductions}")
            if bonus_notes:
                notes_parts.append(f"Bonus: {bonus_notes}")
            combined_notes = "\n\n".join(notes_parts)

            j1_data_by_raw[raw] = {
                "timing_adjustment": timing,
                "eval_bonus": eval_bonus,
                "final_score": final_score,
                "qualification": str(qual).strip() if qual else None,
                "notes": combined_notes,
            }

        # 6. Parse Judge 2 Sheet
        s_j2 = wb["Judge 2 sheet"]
        j2_data_by_raw = {}
        for r in range(7, s_j2.max_row + 1):
            t_name = s_j2.cell(r, 1).value
            if not t_name:
                continue
            raw = str(t_name).strip()

            crit_scores = {}
            crit_comments = {}

            # Criterion mappings: col pairs (score, justification)
            pairs = [
                ("mission_concept", 2, 3),
                ("dashboard", 4, 5),
                ("madar_mission", 6, 7),
                ("tinkercad_circuit", 8, 9),
                ("tinkercad_3d_design", 10, 11),
                ("poster", 12, 13),
                ("recorded_presentation", 14, 15),
                ("consistency", 16, 17),
            ]
            for c_key, s_col, j_col in pairs:
                sc_val = s_j2.cell(r, s_col).value
                if sc_val is not None:
                    try:
                        crit_scores[c_key] = float(sc_val)
                    except (ValueError, TypeError):
                        pass
                comm_val = s_j2.cell(r, j_col).value
                if comm_val:
                    crit_comments[c_key] = str(comm_val).strip()

            report_bonus_val = s_j2.cell(r, 18).value
            report_bonus = float(report_bonus_val) if report_bonus_val is not None else 0.0
            report_comm = s_j2.cell(r, 19).value
            if report_comm:
                crit_comments["report_bonus"] = str(report_comm).strip()

            timing_val = s_j2.cell(r, 21).value
            timing = float(timing_val) if timing_val is not None else 0.0

            final_val = s_j2.cell(r, 23).value
            final_score = float(final_val) if final_val is not None else 0.0

            overall_notes = s_j2.cell(r, 25).value or ""
            favorite = s_j2.cell(r, 26).value
            if favorite:
                overall_notes = f"{overall_notes}\n[Judge Favorite: {favorite}]".strip()

            j2_data_by_raw[raw] = {
                "criterion_scores": crit_scores,
                "criterion_comments": crit_comments,
                "bonus_points": report_bonus,
                "timing_adjustment": timing,
                "final_score": final_score,
                "notes": overall_notes,
            }

        # 7. Parse Final Ranking Sheet (Master Roster & Order)
        s_final = wb["Final Ranking"]
        records = []
        unmatched_count = 0

        print("\n" + "=" * 95)
        print(f"{'Rank':4} | {'Excel Team Name':40} | {'Platform Team':25} | {'Match Type':20}")
        print("=" * 95)

        for r in range(5, 38):
            rank_val = s_final.cell(r, 1).value
            name_val = s_final.cell(r, 2).value
            if not name_val:
                continue
            raw_name = str(name_val).strip()
            rank = int(rank_val)

            j1_score = float(s_final.cell(r, 3).value or 0.0)
            j2_adj = float(s_final.cell(r, 5).value or 0.0)
            comb_pct = float(s_final.cell(r, 7).value or 0.0)
            qual = str(s_final.cell(r, 8).value or "").strip()

            team, match_method = find_team(raw_name)
            if not team:
                unmatched_count += 1
                match_str = "❌ UNMATCHED"
                platform_str = "NOT FOUND"
            else:
                match_str = match_method
                platform_str = f"{team.name} ({team.id.hex[:6]}…)"

            print(f"{rank:4d} | {raw_name:40} | {platform_str:25} | {match_str:20}")

            # Lookup Judge 1 details
            j1_detail = j1_data_by_raw.get(raw_name)
            if not j1_detail:
                # Fallback search by clean name
                j1_detail = next((v for k, v in j1_data_by_raw.items() if clean_team_name(k).lower() == clean_team_name(raw_name).lower()), None)

            # Lookup Judge 2 details
            j2_detail = j2_data_by_raw.get(raw_name)
            if not j2_detail:
                j2_detail = next((v for k, v in j2_data_by_raw.items() if clean_team_name(k).lower() == clean_team_name(raw_name).lower()), None)

            records.append({
                "rank": rank,
                "raw_name": raw_name,
                "team": team,
                "combined_pct": comb_pct,
                "qualification": qual,
                "j1_score": j1_score,
                "j1_detail": j1_detail,
                "j2_score": j2_adj,
                "j2_detail": j2_detail,
            })

        print("=" * 95)
        print(f"\n📊 Summary: {len(records)} teams processed. Matched: {len(records) - unmatched_count}, Unmatched: {unmatched_count}")

        if unmatched_count > 0:
            print(f"\n❌ Aborting: {unmatched_count} team(s) could not be matched automatically.")
            sys.exit(1)

        if not args.commit:
            print("\n🔍 DRY-RUN MODE: No database records were written.")
            print("To commit these scores and team qualifications into the database, re-run with --commit.")
            return

        # 8. Commit to Database
        print("\n💾 Committing scores and qualifications to database...")

        for rec in records:
            team = rec["team"]

            # Update team qualification (only keep official TOP 10 tag)
            team.qualification = "TOP 10" if (rec["qualification"] and "TOP 10" in str(rec["qualification"]).upper()) else None

            # ── Upsert Judge 1 Score ──
            j1_detail = rec["j1_detail"] or {}
            j1_final = rec["j1_score"] or j1_detail.get("final_score", 0.0)
            j1_timing = j1_detail.get("timing_adjustment", 0.0)
            j1_notes = j1_detail.get("notes", "")

            score1 = await db.scalar(
                select(CompetitionScore).where(
                    CompetitionScore.rubric_id == rubric.id,
                    CompetitionScore.team_id == team.id,
                    CompetitionScore.judge_user_id == judge1.id,
                    CompetitionScore.cohort_id == cohort_uuid,
                )
            )
            if score1 is None:
                score1 = CompetitionScore(
                    id=uuid.uuid4(),
                    rubric_id=rubric.id,
                    team_id=team.id,
                    judge_user_id=judge1.id,
                    cohort_id=cohort_uuid,
                    criterion_scores={},
                    criterion_comments={},
                    bonus_points=0.0,
                    timing_adjustment=j1_timing,
                    total_points=j1_final,
                    notes=j1_notes,
                    submitted=True,
                    submitted_at=datetime.now(timezone.utc),
                )
                db.add(score1)
            else:
                score1.timing_adjustment = j1_timing
                score1.total_points = j1_final
                score1.notes = j1_notes
                score1.submitted = True
                score1.submitted_at = datetime.now(timezone.utc)

            # ── Upsert Judge 2 Score ──
            j2_detail = rec["j2_detail"] or {}
            j2_final = rec["j2_score"] or j2_detail.get("final_score", 0.0)
            j2_timing = j2_detail.get("timing_adjustment", 0.0)
            j2_bonus = j2_detail.get("bonus_points", 0.0)
            j2_crit_scores = j2_detail.get("criterion_scores", {})
            j2_crit_comments = j2_detail.get("criterion_comments", {})
            j2_notes = j2_detail.get("notes", "")

            score2 = await db.scalar(
                select(CompetitionScore).where(
                    CompetitionScore.rubric_id == rubric.id,
                    CompetitionScore.team_id == team.id,
                    CompetitionScore.judge_user_id == judge2.id,
                    CompetitionScore.cohort_id == cohort_uuid,
                )
            )
            if score2 is None:
                score2 = CompetitionScore(
                    id=uuid.uuid4(),
                    rubric_id=rubric.id,
                    team_id=team.id,
                    judge_user_id=judge2.id,
                    cohort_id=cohort_uuid,
                    criterion_scores=j2_crit_scores,
                    criterion_comments=j2_crit_comments,
                    bonus_points=j2_bonus,
                    timing_adjustment=j2_timing,
                    total_points=j2_final,
                    notes=j2_notes,
                    submitted=True,
                    submitted_at=datetime.now(timezone.utc),
                )
                db.add(score2)
            else:
                score2.criterion_scores = j2_crit_scores
                score2.criterion_comments = j2_crit_comments
                score2.bonus_points = j2_bonus
                score2.timing_adjustment = j2_timing
                score2.total_points = j2_final
                score2.notes = j2_notes
                score2.submitted = True
                score2.submitted_at = datetime.now(timezone.utc)

        await db.commit()
        print(f"\n✅ SUCCESS! All {len(records)} teams' scores, timing adjustments, comments, and qualifications committed to cohort {cohort_uuid}.")


if __name__ == "__main__":
    asyncio.run(main())
