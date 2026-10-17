"""Bulk-add cohort quiz questions from a filled-in Excel template. Mirrors
`services/instructor_assessment/importer.py`'s parse/generate pair exactly
(header-name column lookup, "collect per-row errors and keep going")
targeting this domain's own column shape — matches the boss-provided
100-question bank's columns directly.
"""

from __future__ import annotations

import io
import re

from pydantic import ValidationError

from app.schemas.lms_cohort_quiz import CohortQuizQuestionIn, CohortQuizQuestionOptionIn

OPTION_LETTERS = "ABCD"
TEMPLATE_HEADERS = [
    "ID", "Topic", "Subtopic", "Difficulty", "Bloom Level", "Time (sec)", "Question Type", "Question",
    "A", "B", "C", "D", "Correct Answer", "Explanation / Rationale", "Source / Basis",
]


def parse_cohort_quiz_question_import(file_bytes: bytes) -> tuple[list[CohortQuizQuestionIn], list[str]]:
    """One row per question. Pure parsing only — no de-dup happens here.
    `ID` isn't a model field at all (each question gets its own real UUID)
    and is never read by this parser; re-running the same file safely
    (after fixing a few error rows) is handled by the caller, via
    `services/lms/cohort_quiz.py::import_questions_deduped`, which skips
    any parsed row whose prompt text already exists in the bank."""
    import openpyxl

    questions: list[CohortQuizQuestionIn] = []
    errors: list[str] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        errors.append(f"Cannot open file: {e}")
        return questions, errors

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        errors.append("The sheet is empty.")
        return questions, errors

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]

    def col(name: str) -> int | None:
        try:
            return headers.index(name.lower())
        except ValueError:
            return None

    topic_col = col("Topic")
    subtopic_col = col("Subtopic")
    difficulty_col = col("Difficulty")
    bloom_col = col("Bloom Level")
    time_col = col("Time (sec)")
    qtype_col = col("Question Type")
    prompt_col = col("Question")
    option_cols = [col(letter) for letter in OPTION_LETTERS]
    correct_col = col("Correct Answer")
    explanation_col = col("Explanation / Rationale")
    source_col = col("Source / Basis")

    if (
        topic_col is None or difficulty_col is None or time_col is None or prompt_col is None
        or correct_col is None or any(c is None for c in option_cols)
    ):
        errors.append(
            "Sheet is missing required columns: Topic, Difficulty, Time (sec), Question, "
            "A-D, Correct Answer."
        )
        return questions, errors

    def cell(row: tuple, idx: int | None):
        return row[idx] if idx is not None and idx < len(row) else None

    def text_or_none(row: tuple, idx: int | None) -> str | None:
        val = cell(row, idx)
        text = str(val).strip() if val is not None else ""
        return text or None

    for row_idx, row in enumerate(rows[1:], start=2):
        prompt = text_or_none(row, prompt_col)
        if not prompt:
            continue  # blank row — skip silently

        topic = text_or_none(row, topic_col)
        if not topic:
            errors.append(f"Row {row_idx}: Topic is required")
            continue

        difficulty_raw = text_or_none(row, difficulty_col)
        difficulty = difficulty_raw.title() if difficulty_raw else None
        if difficulty not in ("Easy", "Medium", "Hard"):
            errors.append(f"Row {row_idx}: Difficulty must be one of Easy, Medium, Hard")
            continue

        time_raw = cell(row, time_col)
        time_str = str(time_raw).strip().lower() if time_raw is not None else ""
        if not time_str or time_str in ("0", "none", "untimed", "no", "n/a"):
            time_limit_seconds = 0
        else:
            try:
                time_limit_seconds = int(float(time_raw))
                if time_limit_seconds < 0:
                    raise ValueError
            except (TypeError, ValueError):
                errors.append(f"Row {row_idx}: Time (sec) must be a whole number (>= 0)")
                continue

        qtype_raw = text_or_none(row, qtype_col)
        is_textarea = bool(qtype_raw and qtype_raw.lower() in ("textarea", "text", "open", "open-ended"))

        if is_textarea:
            try:
                questions.append(CohortQuizQuestionIn(
                    topic=topic,
                    subtopic=text_or_none(row, subtopic_col),
                    difficulty=difficulty,
                    time_limit_seconds=time_limit_seconds,
                    bloom_level=text_or_none(row, bloom_col),
                    question_type="textarea",
                    prompt=prompt,
                    options=[],
                    explanation=text_or_none(row, explanation_col),
                    source_url=text_or_none(row, source_col),
                ))
            except ValidationError as e:
                errors.append(f"Row {row_idx}: {e.errors()[0]['msg']}")
            continue

        options = []
        missing_option = False
        for letter, idx in zip(OPTION_LETTERS, option_cols):
            text = text_or_none(row, idx)
            if not text:
                errors.append(f"Row {row_idx}: Option {letter} is required")
                missing_option = True
                break
            options.append((letter, text))
        if missing_option:
            continue

        correct_raw = cell(row, correct_col)
        correct_str = str(correct_raw).strip().upper() if correct_raw is not None else ""
        is_opinion = correct_str in ("", "NONE", "N/A", "OPINION")

        parts = [p.strip() for p in re.split(r"[,; /]+", correct_str) if p.strip()] if not is_opinion else []
        is_multiselect = bool(
            (qtype_raw and qtype_raw.lower() in ("multiselect", "multi", "multiple_choice", "multiple-choice"))
            or len(parts) > 1
        )

        if not is_opinion:
            invalid = [p for p in parts if p not in OPTION_LETTERS]
            if invalid or not parts:
                errors.append(f"Row {row_idx}: Correct Answer must be one or more of A, B, C, D (or None/blank for opinion-based)")
                continue

        final_qtype = "multiselect" if is_multiselect else "mcq"
        try:
            questions.append(CohortQuizQuestionIn(
                topic=topic,
                subtopic=text_or_none(row, subtopic_col),
                difficulty=difficulty,
                time_limit_seconds=time_limit_seconds,
                bloom_level=text_or_none(row, bloom_col),
                question_type=final_qtype,
                prompt=prompt,
                options=[
                    CohortQuizQuestionOptionIn(text=text, is_correct=False if is_opinion else (letter in parts))
                    for letter, text in options
                ],
                explanation=text_or_none(row, explanation_col),
                source_url=text_or_none(row, source_col),
            ))
        except ValidationError as e:
            errors.append(f"Row {row_idx}: {e.errors()[0]['msg']}")

    return questions, errors


def generate_cohort_quiz_question_template_xlsx() -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Questions"
    for col_idx, header in enumerate(TEMPLATE_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="653f84")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = 22
    example = [
        1, "Introduction to Space & Satellites", "Satellite basics", "Easy", "Remember", 15, "Flash MCQ",
        "Which statement best defines an artificial satellite?",
        "A human-made object placed in orbit", "A natural body orbiting a planet",
        "A vehicle used only during launch", "A station fixed above Earth's surface",
        "A", "Artificial satellites are human-made objects placed into an orbit around a larger body.",
        "https://example.com/source",
    ]
    for col_idx, val in enumerate(example, 1):
        ws.cell(row=2, column=col_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
