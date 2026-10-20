"""Results export (2026-08-27) — one .xlsx with everything ops needs to
review a cohort's judging round off-platform: a Scores sheet (rank,
average, every judge's own score) and a Submissions sheet (every team's
link-only deliverables, who submitted them, and when). Reuses
`judging.leaderboard` and `lms.program.team_submission_items` rather than
re-querying — same numbers the Results tab and judge team-detail view
already show, never a second source of truth for either.
"""

import io
import uuid

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.lms.cohort_quiz import CohortQuizAnswer, CohortQuizAttempt, CohortQuizQuestion
from app.models.sessions.cohort import Cohort
from app.models.sessions.registration import Registration
from app.models.spine.contact import Contact
from app.models.spine.organization import Organization
from app.models.team import Team, TeamMember
from app.models.user import User
from app.services.competitions.judging import leaderboard
from app.services.lms.program import effective_cohort_program_items, team_submission_items

MAX_COLUMN_WIDTH = 60


def _style_header(sheet: Worksheet) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24


def _autosize(sheet: Worksheet) -> None:
    for col_cells in sheet.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        sheet.column_dimensions[col_cells[0].column_letter].width = min(length + 2, MAX_COLUMN_WIDTH)


async def export_results_workbook(db: AsyncSession, *, cohort_id: uuid.UUID) -> bytes:
    board = await leaderboard(db, cohort_id=cohort_id)
    # Stable column order across runs, not insertion order off a set.
    judge_names = sorted({js["judge_name"] for row in board for js in row["judge_scores"]})

    wb = openpyxl.Workbook()
    scores = wb.active
    scores.title = "Scores"
    scores.append(["Rank", "Team", "Avg Score", "Status", "Advanced To", "Score %", "Qualification", *judge_names])
    for row in board:
        by_judge = {js["judge_name"]: js["total_points"] for js in row["judge_scores"]}
        scores.append([
            row["rank"], row["team_name"], row["score"],
            row["status"], row["advanced_to"] or "",
            f"{row['score_percent']}%" if row.get("score_percent") is not None else "",
            row.get("qualification") or "",
            *[by_judge.get(name, "") for name in judge_names],
        ])
    _autosize(scores)

    submissions = wb.create_sheet("Submissions")
    submissions.append(["Team", "Item", "Type", "Status", "Link", "Submitted At", "Submitted By"])
    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort_id).order_by(Team.name))).scalars().all()
    for team in teams:
        for item in await team_submission_items(db, team_id=team.id):
            links = item.get("submitted_links")
            # Multi-field items (2026-09-14, text fields added 2026-09-16)
            # get every labeled field crammed into the one "Link" cell, one
            # per line — a second Link2/Link3 column would need to be as
            # wide as the max fields any item on this cohort ever has,
            # which nothing here tracks up front. A text field's answer
            # goes in the same cell as any link's, since it's still "what
            # this team submitted" either way.
            if links:
                link_text = "\n".join(
                    f"{link['label']}: {link['url'] or link['text']}" if link.get("label") else (link["url"] or link["text"])
                    for link in links
                )
            else:
                # A plain (non-multi) item can itself be a text field
                # (2026-09-16) — same cell either way.
                link_text = item["submitted_url"] or item.get("submitted_text") or ""
            submissions.append([
                team.name, item["title"], item["item_type"], item["status"],
                link_text,
                item["updated_at"].strftime("%Y-%m-%d %H:%M") if item["updated_at"] else "",
                item["updated_by_name"] or "",
            ])
    _autosize(submissions)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_quiz_answers_workbook(db: AsyncSession, *, cohort_id: uuid.UUID) -> bytes:
    """Exports all quiz answers for this cohort as a single .xlsx workbook:
    - 'Quiz Summary' sheet: Every team member's quiz completion status and score
    - 'Quiz Answers' sheet: Every student's answer per question, grouped by Team, Quiz, and Student
    """
    teams = list((await db.execute(
        select(Team).where(Team.cohort_id == cohort_id).order_by(Team.name)
    )).scalars().all())
    team_ids = [t.id for t in teams]

    if team_ids:
        tm_rows = (await db.execute(
            select(TeamMember.team_id, TeamMember.user_id).where(TeamMember.team_id.in_(team_ids))
        )).all()
        members_by_team: dict[uuid.UUID, list[uuid.UUID]] = {tid: [] for tid in team_ids}
        all_user_ids: set[uuid.UUID] = set()
        for tid, uid in tm_rows:
            members_by_team[tid].append(uid)
            all_user_ids.add(uid)
        users = (await db.execute(select(User).where(User.id.in_(all_user_ids)))).scalars().all() if all_user_ids else []
        users_by_id = {u.id: u for u in users}
    else:
        members_by_team = {}
        users_by_id = {}

    resolved = await effective_cohort_program_items(db, cohort_id)
    items = list(resolved[1]) if resolved is not None else []
    quiz_items = [i for i in items if i.item_type == "quiz"]

    wb = openpyxl.Workbook()
    summary_sheet = wb.active
    summary_sheet.title = "Quiz Summary"

    if not quiz_items:
        summary_sheet.append(["Notice", "No quiz items found on this cohort's program checklist."])
        _style_header(summary_sheet)
        _autosize(summary_sheet)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    quiz_item_ids = [i.id for i in quiz_items]
    attempts = (await db.execute(
        select(CohortQuizAttempt).where(
            CohortQuizAttempt.program_item_id.in_(quiz_item_ids),
            (CohortQuizAttempt.cohort_id == cohort_id) | (CohortQuizAttempt.user_id.in_(list(users_by_id.keys()))),
        )
    )).scalars().all() if users_by_id else []

    attempts_by_user_item: dict[tuple[uuid.UUID, uuid.UUID], CohortQuizAttempt] = {
        (a.user_id, a.program_item_id): a for a in attempts
    }

    submitted_attempts = [a for a in attempts if a.status == "submitted"]
    submitted_attempt_ids = [a.id for a in submitted_attempts]
    answers = (await db.execute(
        select(CohortQuizAnswer).where(CohortQuizAnswer.attempt_id.in_(submitted_attempt_ids))
    )).scalars().all() if submitted_attempt_ids else []
    answers_by_attempt_qid = {(a.attempt_id, a.question_id): a for a in answers}

    all_qids: set[uuid.UUID] = set()
    for a in submitted_attempts:
        if a.question_order:
            for entry in a.question_order:
                try:
                    all_qids.add(uuid.UUID(entry["question_id"]))
                except (ValueError, KeyError):
                    pass
    questions = (await db.execute(
        select(CohortQuizQuestion).where(CohortQuizQuestion.id.in_(all_qids))
    )).scalars().all() if all_qids else []
    questions_by_id = {q.id: q for q in questions}

    # Populate Sheet 1: Quiz Summary
    summary_sheet.append([
        "Team", "Team ID", "Student Name", "Email", "Quiz", "Status",
        "Score", "Total", "Percentage", "Submitted At",
    ])
    for team in teams:
        member_ids = members_by_team.get(team.id, [])
        for q_item in quiz_items:
            for uid in member_ids:
                u = users_by_id.get(uid)
                att = attempts_by_user_item.get((uid, q_item.id))
                status_label = "Not Started"
                score_val = ""
                total_val = ""
                pct_val = ""
                sub_at_val = ""
                if att:
                    if att.status == "submitted":
                        status_label = "Submitted"
                        score_val = att.score if att.score is not None else 0
                        total_val = att.total if att.total is not None else 0
                        pct_val = f"{round(100 * score_val / total_val)}%" if total_val else "0%"
                        sub_at_val = att.submitted_at.strftime("%Y-%m-%d %H:%M") if att.submitted_at else ""
                    elif att.status == "in_progress":
                        status_label = "In Progress"
                        total_val = att.total if att.total is not None else ""

                summary_sheet.append([
                    team.name, team.roster_team_id or "", u.full_name if u else "Student",
                    u.email if u else "", q_item.title, status_label,
                    score_val, total_val, pct_val, sub_at_val,
                ])
    _style_header(summary_sheet)
    _autosize(summary_sheet)

    # Populate Sheet 2: Quiz Answers
    answers_sheet = wb.create_sheet("Quiz Answers")
    answers_sheet.append([
        "Team", "Student Name", "Quiz", "Q#", "Topic", "Question Prompt",
        "Student Answer", "Correct Answer", "Result", "Time (s)", "Timed Out", "Explanation",
    ])

    for team in teams:
        member_ids = members_by_team.get(team.id, [])
        for q_item in quiz_items:
            for uid in member_ids:
                u = users_by_id.get(uid)
                att = attempts_by_user_item.get((uid, q_item.id))
                if not att or att.status != "submitted" or not att.question_order:
                    continue

                for pos, entry in enumerate(att.question_order, start=1):
                    try:
                        qid = uuid.UUID(entry["question_id"])
                    except (ValueError, KeyError):
                        continue
                    q = questions_by_id.get(qid)
                    if not q:
                        continue
                    ans = answers_by_attempt_qid.get((att.id, qid))
                    if not ans:
                        continue

                    student_ans_text = ""
                    if ans.selected_option_index is not None and 0 <= ans.selected_option_index < len(q.options):
                        student_ans_text = q.options[ans.selected_option_index].get("text", "")
                    elif ans.timed_out:
                        student_ans_text = "(Timed out)"
                    else:
                        student_ans_text = "(Unanswered)"

                    correct_opt = next((o.get("text", "") for o in q.options if o.get("is_correct")), "")

                    answers_sheet.append([
                        team.name, u.full_name if u else "Student", q_item.title,
                        pos, q.topic, q.prompt,
                        student_ans_text, correct_opt,
                        "Correct" if ans.is_correct else "Incorrect",
                        ans.time_taken_seconds if ans.time_taken_seconds is not None else "",
                        "Yes" if ans.timed_out else "No",
                        q.explanation or "",
                    ])

    _style_header(answers_sheet)
    _autosize(answers_sheet)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_teams_workbook(db: AsyncSession, *, cohort_id: uuid.UUID) -> bytes:
    """Exports all teams and their students for a cohort:
    - Team ID
    - Team Name
    - Student Name
    - Student Email
    - Phone Number
    - Organization
    Grouped with members of the same team listed under each other.
    """
    teams = list((await db.execute(
        select(Team).where(Team.cohort_id == cohort_id).order_by(Team.name)
    )).scalars().all())
    team_ids = [t.id for t in teams]

    if team_ids:
        tm_rows = (await db.execute(
            select(TeamMember.team_id, TeamMember.user_id).where(TeamMember.team_id.in_(team_ids))
        )).all()
        members_by_team: dict[uuid.UUID, list[uuid.UUID]] = {tid: [] for tid in team_ids}
        all_user_ids: set[uuid.UUID] = set()
        for tid, uid in tm_rows:
            members_by_team[tid].append(uid)
            all_user_ids.add(uid)
        users = (await db.execute(select(User).where(User.id.in_(all_user_ids)))).scalars().all() if all_user_ids else []
        users_by_id = {u.id: u for u in users}
    else:
        members_by_team = {}
        users_by_id = {}

    cohort = await db.get(Cohort, cohort_id)
    default_org_name = ""
    if cohort and cohort.organization_id:
        cohort_org = await db.get(Organization, cohort.organization_id)
        if cohort_org:
            default_org_name = cohort_org.name_latin or cohort_org.name_native or ""

    reg_stmt = (
        select(
            Contact.email,
            Contact.primary_phone_e164,
            Organization.name_latin.label("org_latin"),
            Organization.name_native.label("org_native"),
        )
        .select_from(Registration)
        .join(Contact, Contact.id == Registration.contact_id)
        .outerjoin(Organization, Organization.id == Contact.organization_id)
        .where(Registration.cohort_id == cohort_id)
    )
    reg_rows = (await db.execute(reg_stmt)).all()
    contact_info_by_email: dict[str, dict[str, str]] = {}
    for r in reg_rows:
        if r.email:
            org = r.org_latin or r.org_native or default_org_name
            contact_info_by_email[r.email.strip().lower()] = {
                "phone": r.primary_phone_e164 or "",
                "organization": org or "",
            }

    user_emails = [u.email.strip().lower() for u in users_by_id.values() if u.email]
    missing_emails = [e for e in user_emails if e not in contact_info_by_email]
    if missing_emails:
        c_stmt = (
            select(
                Contact.email,
                Contact.primary_phone_e164,
                Organization.name_latin.label("org_latin"),
                Organization.name_native.label("org_native"),
            )
            .select_from(Contact)
            .outerjoin(Organization, Organization.id == Contact.organization_id)
            .where(func.lower(Contact.email).in_(missing_emails))
        )
        for r in (await db.execute(c_stmt)).all():
            if r.email:
                org = r.org_latin or r.org_native or default_org_name
                contact_info_by_email[r.email.strip().lower()] = {
                    "phone": r.primary_phone_e164 or "",
                    "organization": org or "",
                }

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Teams & Students"
    sheet.append([
        "Team ID",
        "Team Name",
        "Student Name",
        "Student Email",
        "Phone Number",
        "Organization",
    ])

    for team in teams:
        uids = members_by_team.get(team.id, [])
        if not uids:
            sheet.append([
                str(team.id),
                team.name,
                "(No members)",
                "",
                "",
                default_org_name,
            ])
        else:
            team_users = [users_by_id[uid] for uid in uids if uid in users_by_id]
            team_users.sort(key=lambda u: (u.full_name or "").lower())
            for u in team_users:
                em_key = (u.email or "").strip().lower()
                c_info = contact_info_by_email.get(em_key, {})
                phone = c_info.get("phone") or u.phone or ""
                org = c_info.get("organization") or default_org_name
                sheet.append([
                    str(team.id),
                    team.name,
                    u.full_name or "",
                    u.email or "",
                    phone,
                    org,
                ])

    _style_header(sheet)
    _autosize(sheet)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

