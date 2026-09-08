"""Bulk-add instructor assessment questions from a filled-in Excel template.

Mirrors `services/games/importer.py`'s parse/generate pair exactly (same
openpyxl usage, same "collect per-row errors and keep going" posture) —
this targets `AssessmentQuestionIn` and the real question bank's own column
names instead of the Kahoot game's. A bad row is skipped and reported, not
fatal: the good rows in the same file still get created.
"""

from __future__ import annotations

import io

from pydantic import ValidationError

from app.schemas.instructor_assessment import AssessmentQuestionIn, AssessmentQuestionOptionIn

OPTION_LETTERS = "ABCD"
TEMPLATE_HEADERS = [
    "Category", "Scenario / Question", "Option A", "Option B", "Option C", "Option D",
    "Correct Answer", "Rationale / Instructor Standard", "Critical Scenario?",
]


def parse_assessment_question_import(file_bytes: bytes) -> tuple[list[AssessmentQuestionIn], list[str]]:
    """One row per question, matching the real 71-question bank's own
    column shape. All four options are required (the real bank always has
    exactly four); "Critical Scenario?" accepts Yes/No (case-insensitive),
    blank = No."""
    import openpyxl

    questions: list[AssessmentQuestionIn] = []
    errors: list[str] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        errors.append(f"Cannot open file: {e}")
        return questions, errors

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        errors.append("The sheet is empty.")
        return questions, errors

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]

    def col(name: str) -> int | None:
        try:
            return headers.index(name.lower())
        except ValueError:
            return None

    category_col = col("Category")
    prompt_col = col("Scenario / Question")
    option_cols = [col(f"Option {letter}") for letter in OPTION_LETTERS]
    correct_col = col("Correct Answer")
    rationale_col = col("Rationale / Instructor Standard")
    critical_col = col("Critical Scenario?")

    if category_col is None or prompt_col is None or correct_col is None or any(c is None for c in option_cols):
        errors.append(
            "Sheet is missing required columns: Category, Scenario / Question, "
            "Option A-D, Correct Answer."
        )
        return questions, errors

    def cell(row: tuple, idx: int | None):
        return row[idx] if idx is not None and idx < len(row) else None

    for row_idx, row in enumerate(rows[1:], start=2):
        prompt_raw = cell(row, prompt_col)
        prompt = str(prompt_raw).strip() if prompt_raw is not None else ""
        if not prompt:
            continue  # blank row — skip silently

        category_raw = cell(row, category_col)
        category = str(category_raw).strip() if category_raw is not None else ""
        if not category:
            errors.append(f"Row {row_idx}: Category is required")
            continue

        options = []
        for letter, idx in zip(OPTION_LETTERS, option_cols):
            val = cell(row, idx)
            text = str(val).strip() if val is not None else ""
            if not text:
                errors.append(f"Row {row_idx}: Option {letter} is required")
                break
            options.append((letter, text))
        else:
            correct_raw = cell(row, correct_col)
            correct_letter = str(correct_raw).strip().upper() if correct_raw is not None else ""
            if correct_letter not in OPTION_LETTERS:
                errors.append(f"Row {row_idx}: Correct Answer must be one of A, B, C, D")
                continue

            rationale_raw = cell(row, rationale_col)
            rationale = str(rationale_raw).strip() if rationale_raw not in (None, "") else None

            critical_raw = cell(row, critical_col)
            critical = str(critical_raw).strip().lower() in ("yes", "y", "true", "1") if critical_raw not in (None, "") else False

            try:
                questions.append(AssessmentQuestionIn(
                    category=category,
                    prompt=prompt,
                    options=[
                        AssessmentQuestionOptionIn(text=text, is_correct=letter == correct_letter)
                        for letter, text in options
                    ],
                    critical=critical,
                    rationale=rationale,
                ))
            except ValidationError as e:
                errors.append(f"Row {row_idx}: {e.errors()[0]['msg']}")
        # else: the `break` above (a missing option) already logged an error;
        # nothing left to do for this row.

    return questions, errors


def generate_assessment_question_template_xlsx() -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Questions"
    for col_idx, header in enumerate(TEMPLATE_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="653f84")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = 26
    example = [
        "Safety & PPE",
        "Students are excited to begin integration, but several have not put on the required PPE. What should you do?",
        "PPE is optional.",
        "Allow them to start and remind them later.",
        "Only ask once and then ignore it.",
        "Pause practical work until the required PPE is worn correctly.",
        "D",
        "Safety requirements take priority over schedule or excitement.",
        "Yes",
    ]
    for col_idx, val in enumerate(example, 1):
        ws.cell(row=2, column=col_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
