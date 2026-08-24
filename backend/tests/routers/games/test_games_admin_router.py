"""Live Games Phase 2C, 8-3 — facilitator authoring: `/games/admin/*`.
Redis-free (uses the `client` fixture).
"""

import io
import uuid

import openpyxl
import pytest
from fastapi import status as http_status

from app.core.security import create_access_token
from app.models.games.game import Game, GameQuestion
from app.models.user import User
from app.services.games.importer import TEMPLATE_HEADERS


async def _user(db, *, roles=None) -> User:
    user = User(
        id=uuid.uuid4(), full_name="Games Admin User", email=f"ga-{uuid.uuid4().hex[:8]}@example.com",
        password_hash="x", roles=list(roles) if roles else ["operations"], status="active",
    )
    db.add(user)
    await db.flush()
    return user


def _headers(user: User) -> dict:
    return {"Authorization": f"Bearer {create_access_token(user.id, user.role_values)}"}


async def _game(db, *, author) -> Game:
    game = Game(id=uuid.uuid4(), title="Transfer Orbits", created_by=author.id)
    db.add(game)
    await db.flush()
    return game


def _mc(prompt="Which burn raises apoapsis?", correct=1):
    return {
        "prompt": prompt,
        "options": [
            {"text": "A", "is_correct": correct == 0},
            {"text": "B", "is_correct": correct == 1},
            {"text": "C", "is_correct": correct == 2},
        ],
    }


def _build_question_workbook(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(TEMPLATE_HEADERS)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_admin_routes_require_content_role(db, client):
    ops = await _user(db, roles=["operations"])
    student = await _user(db, roles=["student"])
    await db.commit()

    ok = await client.get("/games/admin", headers=_headers(ops))
    assert ok.status_code == 200

    forbidden = await client.get("/games/admin", headers=_headers(student))
    assert forbidden.status_code == http_status.HTTP_403_FORBIDDEN


@pytest.mark.asyncio
async def test_create_game_with_defaults(db, client):
    ops = await _user(db)
    await db.commit()

    resp = await client.post("/games/admin", headers=_headers(ops), json={"title": "Transfer Orbits"})
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["default_time_limit_seconds"] == 20
    assert body["default_floor_pct"] == 25
    assert body["default_blackout_count"] == 3
    assert body["question_count"] == 0


@pytest.mark.asyncio
async def test_create_question_rejects_zero_or_multiple_correct_options(db, client):
    ops = await _user(db)
    game = await _game(db, author=ops)
    await db.commit()

    zero_correct = await client.post(
        f"/games/admin/{game.id}/questions", headers=_headers(ops),
        json={"prompt": "X?", "options": [{"text": "A"}, {"text": "B"}]},
    )
    assert zero_correct.status_code == http_status.HTTP_422_UNPROCESSABLE_ENTITY

    two_correct = await client.post(
        f"/games/admin/{game.id}/questions", headers=_headers(ops),
        json={"prompt": "X?", "options": [{"text": "A", "is_correct": True}, {"text": "B", "is_correct": True}]},
    )
    assert two_correct.status_code == http_status.HTTP_422_UNPROCESSABLE_ENTITY


@pytest.mark.asyncio
async def test_create_question_defaults_to_normal_points_auto_appends_position(db, client):
    ops = await _user(db)
    game = await _game(db, author=ops)
    await db.commit()

    first = await client.post(f"/games/admin/{game.id}/questions", headers=_headers(ops), json=_mc())
    assert first.status_code == 201, first.text
    assert first.json()["position"] == 1
    assert first.json()["points_mode"] == "normal"
    assert first.json()["max_points"] == 100

    second = await client.post(
        f"/games/admin/{game.id}/questions", headers=_headers(ops),
        json={**_mc(prompt="Second?"), "points_mode": "double"},
    )
    assert second.status_code == 201, second.text
    assert second.json()["position"] == 2
    assert second.json()["max_points"] == 200


@pytest.mark.asyncio
async def test_get_game_lists_questions_in_position_order(db, client):
    ops = await _user(db)
    game = await _game(db, author=ops)
    db.add(GameQuestion(id=uuid.uuid4(), game_id=game.id, position=1, prompt="Q1", options=[
        {"text": "A", "is_correct": True}, {"text": "B", "is_correct": False},
    ]))
    db.add(GameQuestion(id=uuid.uuid4(), game_id=game.id, position=2, prompt="Q2", options=[
        {"text": "A", "is_correct": True}, {"text": "B", "is_correct": False},
    ]))
    await db.commit()

    resp = await client.get(f"/games/admin/{game.id}", headers=_headers(ops))
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["question_count"] == 2
    assert [q["prompt"] for q in body["questions"]] == ["Q1", "Q2"]


@pytest.mark.asyncio
async def test_duplicate_question_appends_a_copy(db, client):
    ops = await _user(db)
    game = await _game(db, author=ops)
    await db.commit()
    created = await client.post(f"/games/admin/{game.id}/questions", headers=_headers(ops), json=_mc())
    question_id = created.json()["id"]

    dup = await client.post(f"/games/admin/questions/{question_id}/duplicate", headers=_headers(ops))
    assert dup.status_code == 201, dup.text
    assert dup.json()["position"] == 2
    assert dup.json()["prompt"] == created.json()["prompt"]

    detail = await client.get(f"/games/admin/{game.id}", headers=_headers(ops))
    assert detail.json()["question_count"] == 2


@pytest.mark.asyncio
async def test_reorder_questions_requires_the_full_set_exactly_once(db, client):
    ops = await _user(db)
    game = await _game(db, author=ops)
    await db.commit()
    a = (await client.post(f"/games/admin/{game.id}/questions", headers=_headers(ops), json=_mc(prompt="A"))).json()
    b = (await client.post(f"/games/admin/{game.id}/questions", headers=_headers(ops), json=_mc(prompt="B"))).json()

    bad = await client.post(
        f"/games/admin/{game.id}/questions/reorder", headers=_headers(ops), json={"question_ids": [a["id"]]},
    )
    assert bad.status_code == http_status.HTTP_400_BAD_REQUEST

    good = await client.post(
        f"/games/admin/{game.id}/questions/reorder", headers=_headers(ops),
        json={"question_ids": [b["id"], a["id"]]},
    )
    assert good.status_code == 200, good.text
    assert [q["prompt"] for q in good.json()] == ["B", "A"]


@pytest.mark.asyncio
async def test_delete_question_and_delete_game(db, client):
    ops = await _user(db)
    game = await _game(db, author=ops)
    await db.commit()
    created = await client.post(f"/games/admin/{game.id}/questions", headers=_headers(ops), json=_mc())
    question_id = created.json()["id"]

    deleted = await client.delete(f"/games/admin/questions/{question_id}", headers=_headers(ops))
    assert deleted.status_code == http_status.HTTP_204_NO_CONTENT

    detail = await client.get(f"/games/admin/{game.id}", headers=_headers(ops))
    assert detail.json()["question_count"] == 0

    deleted_game = await client.delete(f"/games/admin/{game.id}", headers=_headers(ops))
    assert deleted_game.status_code == http_status.HTTP_204_NO_CONTENT
    gone = await client.get(f"/games/admin/{game.id}", headers=_headers(ops))
    assert gone.status_code == http_status.HTTP_404_NOT_FOUND


@pytest.mark.asyncio
async def test_download_question_import_template_returns_an_xlsx(db, client):
    ops = await _user(db)
    await db.commit()

    resp = await client.get("/games/admin/questions/import/template", headers=_headers(ops))
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    header = [c.value for c in next(wb.active.iter_rows(max_row=1))]
    assert header == TEMPLATE_HEADERS


@pytest.mark.asyncio
async def test_import_questions_creates_valid_rows_and_reports_bad_ones(db, client):
    ops = await _user(db)
    game = await _game(db, author=ops)
    await db.commit()
    # one question already exists — imported rows must append after it
    existing = await client.post(f"/games/admin/{game.id}/questions", headers=_headers(ops), json=_mc())
    assert existing.status_code == 201, existing.text

    file_bytes = _build_question_workbook([
        ["Which burn raises apoapsis?", "Prograde at periapsis", "Retrograde at periapsis", "Radial", None, "A", 15, "double"],
        ["Second question?", "Yes", "No", None, None, "B", None, None],
        ["Bad correct letter", "Yes", "No", None, None, "Z", None, None],
        [None, None, None, None, None, None, None, None],  # blank row — skipped silently
    ])

    resp = await client.post(
        f"/games/admin/{game.id}/questions/import", headers=_headers(ops),
        files={"file": ("questions.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["created"] == 2
    assert len(body["errors"]) == 1 and "Row 4" in body["errors"][0]
    assert [q["position"] for q in body["questions"]] == [2, 3]
    assert body["questions"][0]["points_mode"] == "double"
    assert body["questions"][0]["max_points"] == 200
    assert body["questions"][1]["time_limit_seconds"] is None
    assert body["questions"][1]["points_mode"] == "normal"

    detail = await client.get(f"/games/admin/{game.id}", headers=_headers(ops))
    assert detail.json()["question_count"] == 3


@pytest.mark.asyncio
async def test_import_questions_requires_content_role(db, client):
    student = await _user(db, roles=["student"])
    ops = await _user(db)
    game = await _game(db, author=ops)
    await db.commit()

    file_bytes = _build_question_workbook([])
    resp = await client.post(
        f"/games/admin/{game.id}/questions/import", headers=_headers(student),
        files={"file": ("questions.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == http_status.HTTP_403_FORBIDDEN
