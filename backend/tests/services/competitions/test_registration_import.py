"""Competition registration import (2026-08-24, team_id grouping + deferred
naming 2026-09-14) — idempotent on email, teams grouped by team_id (not
name — the roster's own external team code), a bad row errors without
failing the batch.
"""

import io
import uuid

import openpyxl
import pytest
from sqlalchemy import select

from app.models.sessions.cohort import Cohort
from app.models.sessions.program import Program
from app.models.sessions.registration import Registration
from app.models.spine.contact import Contact
from app.models.team import Team
from app.models.user import User
from app.services.competitions.registration_import import (
    REQUIRED_COLUMNS, assign_existing_students, import_registrations, parse_registration_sheet, summarize,
    summarize_assign_results,
)
from app.services.teams import create_team, team_member_ids


async def _make_cohort(db) -> Cohort:
    program = Program(
        id=uuid.uuid4(), code=f"REG-{uuid.uuid4().hex[:8]}", name="Ru'ya Space Race",
        program_type="competition", pricing_model="free", active=True,
    )
    db.add(program)
    await db.flush()
    cohort = Cohort(id=uuid.uuid4(), program_id=program.id, name="Qualifier", status="registration_open")
    db.add(cohort)
    await db.flush()
    return cohort


def _workbook(rows: list[dict]) -> bytes:
    """Column set is the union of every row's own keys, in first-seen
    order — lets each test only spell out the fields it actually cares
    about (real sheets vary in which optional columns they carry too)."""
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(c) for c in columns])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_parse_requires_columns():
    from fastapi import HTTPException
    wb = openpyxl.Workbook()
    wb.active.append(["team_id", "member_name"])  # missing email
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(HTTPException):
        parse_registration_sheet(buf.getvalue())
    assert "email" in REQUIRED_COLUMNS
    assert "team_id" in REQUIRED_COLUMNS
    assert "team_name" not in REQUIRED_COLUMNS  # optional now — see module docstring


@pytest.mark.asyncio
async def test_import_creates_teams_and_student_accounts(db):
    cohort = await _make_cohort(db)
    await db.commit()

    file_bytes = _workbook([
        {"team_id": "T1", "team_name": "Falcon Eye", "member_name": "Amir Salem", "email": f"amir-{uuid.uuid4().hex[:6]}@example.com"},
        {"team_id": "T1", "team_name": "Falcon Eye", "member_name": "Lina Nasser", "email": f"lina-{uuid.uuid4().hex[:6]}@example.com"},
        {"team_id": "T2", "team_name": "Silent Comet", "member_name": "Zaid Omar", "email": f"zaid-{uuid.uuid4().hex[:6]}@example.com"},
    ])
    rows = parse_registration_sheet(file_bytes)
    results = await import_registrations(db, cohort_id=cohort.id, rows=rows)
    await db.commit()

    summary = summarize(results)
    assert summary == {"total": 3, "created": 3, "matched": 0, "error": 0}

    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert {t.name for t in teams} == {"Falcon Eye", "Silent Comet"}
    assert {t.roster_team_id for t in teams} == {"T1", "T2"}
    falcon = next(t for t in teams if t.name == "Falcon Eye")
    assert len(await team_member_ids(db, team_id=falcon.id)) == 2

    users = (await db.execute(select(User).where(User.roles.any("student")))).scalars().all()
    assert len(users) >= 3


@pytest.mark.asyncio
async def test_import_falls_back_to_the_team_id_when_team_name_is_blank(db):
    """The roster's whole point (2026-09-14): a team can be onboarded with
    only its id, no name yet — students set the real one themselves later.
    `name` is seeded to `team_id` as a literal placeholder rather than left
    NULL — see app/models/team.py's module docstring."""
    cohort = await _make_cohort(db)
    await db.commit()

    file_bytes = _workbook([
        {"team_id": "T-014", "member_name": "Amir Salem", "email": f"amir-{uuid.uuid4().hex[:6]}@example.com"},
    ])
    results = await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()

    assert summarize(results) == {"total": 1, "created": 1, "matched": 0, "error": 0}
    team = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalar_one()
    assert team.roster_team_id == "T-014"
    assert team.name == "T-014"  # still unnamed — name == roster_team_id


@pytest.mark.asyncio
async def test_reimport_never_clobbers_a_name_the_team_already_set(db):
    """A team that renamed itself (or that ops fixed) between two imports of
    the same sheet must keep that name — only ever apply the sheet's name
    while the team is still at its own placeholder."""
    cohort = await _make_cohort(db)
    await db.commit()

    email = f"amir-{uuid.uuid4().hex[:6]}@example.com"
    file_bytes = _workbook([{"team_id": "T-014", "member_name": "Amir Salem", "email": email}])
    await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()

    team = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalar_one()
    assert team.name == "T-014"
    team.name = "Falcon Eye"  # simulates the team's own rename
    await db.commit()

    # Re-import the exact same sheet (still no team_name column at all).
    results = await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()
    assert summarize(results) == {"total": 1, "created": 0, "matched": 1, "error": 0}

    await db.refresh(team)
    assert team.name == "Falcon Eye"

    # Now the sheet DOES carry a name for this team_id — still must not
    # overwrite the real one already set.
    named_file = _workbook([{"team_id": "T-014", "team_name": "Something Else", "member_name": "Amir Salem", "email": email}])
    await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(named_file))
    await db.commit()
    await db.refresh(team)
    assert team.name == "Falcon Eye"


@pytest.mark.asyncio
async def test_import_flags_a_team_name_collision_without_failing_the_batch(db):
    """Case-insensitive (models/team.py) — an existing real team name
    collides with another team's fallback-to-team_id placeholder just as
    easily as with an explicit team_name, so this covers both directions."""
    from app.core.security import get_password_hash

    cohort = await _make_cohort(db)
    creator = User(
        id=uuid.uuid4(), full_name="Team Creator", email=f"creator-{uuid.uuid4().hex[:6]}@example.com",
        password_hash=get_password_hash("dummy"), roles=["student"], status="active",
    )
    db.add(creator)
    await db.flush()
    existing = await create_team(db, name="Falcon", created_by=creator.id, cohort_id=cohort.id)
    await db.commit()

    file_bytes = _workbook([
        {"team_id": "FALCON", "member_name": "Amir Salem", "email": f"amir-{uuid.uuid4().hex[:6]}@example.com"},
    ])
    results = await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()

    summary = summarize(results)
    assert summary == {"total": 1, "created": 0, "matched": 0, "error": 1}
    assert "already exists" in results[0].reason

    # The other team is untouched, and no second team got created.
    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert len(teams) == 1
    assert teams[0].id == existing.id


@pytest.mark.asyncio
async def test_import_captures_nationality_on_the_contact(db):
    cohort = await _make_cohort(db)
    await db.commit()
    email = f"amir-{uuid.uuid4().hex[:6]}@example.com"
    file_bytes = _workbook([
        {"team_id": "T1", "member_name": "Amir Salem", "email": email, "nationality": "Emirati"},
    ])
    await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()

    contact = (await db.execute(select(Contact).where(Contact.email == email))).scalar_one()
    assert contact.nationality == "Emirati"


@pytest.mark.asyncio
async def test_import_captures_gender_on_the_contact_and_backfills_on_reimport(db):
    from app.services.competitions.registration_import import preview_registration_sheet

    cohort = await _make_cohort(db)
    await db.commit()
    email_f = f"lina-{uuid.uuid4().hex[:6]}@example.com"
    email_m = f"amir-{uuid.uuid4().hex[:6]}@example.com"

    # First import: Lina has Gender="Female", Amir has no Gender yet
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Team ID", "Student Name", "Email", "Gender"])
    sheet.append(["T1", "Lina Nasser", email_f, "Female"])
    sheet.append(["T1", "Amir Salem", email_m, None])
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    preview = preview_registration_sheet(file_bytes)
    assert [r.gender for r in preview] == ["Female", None]

    await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()

    contact_f = (await db.execute(select(Contact).where(Contact.email == email_f))).scalar_one()
    contact_m = (await db.execute(select(Contact).where(Contact.email == email_m))).scalar_one()
    assert contact_f.gender == "Female"
    assert contact_m.gender is None

    # Re-import with Amir's Gender filled in — should gap-fill on existing cohort contact
    wb2 = openpyxl.Workbook()
    sheet2 = wb2.active
    sheet2.append(["Team ID", "Student Name", "Email", "Gender"])
    sheet2.append(["T1", "Lina Nasser", email_f, "Female"])
    sheet2.append(["T1", "Amir Salem", email_m, "Male"])
    buf2 = io.BytesIO()
    wb2.save(buf2)

    await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(buf2.getvalue()))
    await db.commit()
    await db.refresh(contact_m)
    assert contact_m.gender == "Male"


@pytest.mark.asyncio
async def test_import_is_idempotent_on_email(db):
    cohort = await _make_cohort(db)
    await db.commit()
    email = f"repeat-{uuid.uuid4().hex[:6]}@example.com"
    file_bytes = _workbook([{"team_id": "T1", "team_name": "Falcon Eye", "member_name": "Amir Salem", "email": email}])

    first = await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()
    assert summarize(first)["created"] == 1

    second = await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()
    assert summarize(second) == {"total": 1, "created": 0, "matched": 1, "error": 0}

    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert len(teams) == 1


@pytest.mark.asyncio
async def test_reimport_reuses_an_already_registered_contact_even_without_a_matching_alias(db):
    """Production bug (2026-09-15): an email got re-uploaded onto a cohort
    where it was already fully registered and on a team, but landed on a
    *second*, alias-less Contact instead of being recognized as the same
    person — `resolve_or_create_contact`'s alias matching never got a
    chance to find the first one. Simulated here by creating the first
    Contact/User/Registration/team-membership directly (bypassing
    `resolve_or_create_contact` entirely, exactly like the corrupted
    production state), then re-importing the same email: the row must
    reuse that existing contact via `_find_existing_contact_by_email`
    (contacts already registered in this cohort, by email) rather than
    minting a second Contact and a second Registration."""
    from app.core.security import get_password_hash

    from app.services.sessions.registration import register
    from app.services.teams import create_team

    cohort = await _make_cohort(db)
    email = f"dup-{uuid.uuid4().hex[:6]}@example.com"

    contact = Contact(id=uuid.uuid4(), full_name="Mohammad Sultan", email=email, contact_roles=["student"])
    db.add(contact)
    await db.flush()
    user = User(
        id=uuid.uuid4(), full_name="Mohammad Sultan", email=email, password_hash=get_password_hash("dummy"),
        roles=["student"], status="active", contact_id=contact.id, must_change_password=False,
    )
    db.add(user)
    await db.flush()
    registration = await register(
        db, contact_id=contact.id, cohort_id=cohort.id, registered_via="import", payment_status="waived",
    )
    team = await create_team(
        db, name="Falcon Eye", created_by=user.id, cohort_id=cohort.id, member_ids=[user.id], roster_team_id="T1",
    )
    await db.commit()

    file_bytes = _workbook([
        {"team_id": "T1", "team_name": "Falcon Eye", "member_name": "Mohammad Sultan", "email": email},
    ])
    results = await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()

    assert summarize(results) == {"total": 1, "created": 0, "matched": 1, "error": 0}
    assert results[0].registration_id == registration.id

    contacts = (await db.execute(select(Contact).where(Contact.email == email))).scalars().all()
    assert len(contacts) == 1, "a re-upload of an already-registered email must not mint a second Contact"

    registrations = (
        await db.execute(select(Registration).where(Registration.cohort_id == cohort.id))
    ).scalars().all()
    assert len(registrations) == 1, "a re-upload of an already-registered email must not create a second Registration"

    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert len(teams) == 1
    members = await team_member_ids(db, team_id=team.id)
    assert members == [user.id]


@pytest.mark.asyncio
async def test_a_row_missing_a_required_field_errors_without_failing_the_batch(db):
    cohort = await _make_cohort(db)
    await db.commit()
    file_bytes = _workbook([
        {"team_id": "T1", "member_name": "", "email": f"noname-{uuid.uuid4().hex[:6]}@example.com"},
        {"team_id": "T1", "member_name": "Lina Nasser", "email": f"lina2-{uuid.uuid4().hex[:6]}@example.com"},
    ])
    results = await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()

    summary = summarize(results)
    assert summary["error"] == 1 and summary["created"] == 1


@pytest.mark.asyncio
async def test_a_row_missing_team_id_errors_even_with_a_team_name_present(db):
    """team_id, not team_name, is the required grouping key now — a row with
    a name but a blank id still errors (the column exists, so this is the
    per-row blocking-issue path, not the missing-column 422 above it)."""
    cohort = await _make_cohort(db)
    await db.commit()
    file_bytes = _workbook([
        {"team_id": None, "team_name": "Falcon Eye", "member_name": "Amir Salem", "email": f"amir-{uuid.uuid4().hex[:6]}@example.com"},
        {"team_id": "T1", "team_name": "Falcon Eye", "member_name": "Lina Nasser", "email": f"lina-{uuid.uuid4().hex[:6]}@example.com"},
    ])
    results = await import_registrations(db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes))
    await db.commit()
    summary = summarize(results)
    assert summary["error"] == 1 and summary["created"] == 1
    error_row = next(r for r in results if r.disposition == "error")
    assert "Team ID" in error_row.reason


@pytest.mark.asyncio
async def test_import_accepts_the_ops_selection_tracker_headers(db):
    """Ops's real "Eligible"/"Final Selection" tracker sheet uses "Student
    Name" and "School / University", not this module's canonical
    `member_name`/`school` — see `_HEADER_ALIASES`."""
    cohort = await _make_cohort(db)
    await db.commit()

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Team Name", "Team ID", "Student Name", "Email", "Phone", "School / University"])
    sheet.append([
        "Falcon Eye", "T1", "Amir Salem", f"amir-{uuid.uuid4().hex[:6]}@example.com", "0501234567", "SABIS",
    ])
    buf = io.BytesIO()
    wb.save(buf)

    rows = parse_registration_sheet(buf.getvalue())
    results = await import_registrations(db, cohort_id=cohort.id, rows=rows)
    await db.commit()

    assert summarize(results) == {"total": 1, "created": 1, "matched": 0, "error": 0}
    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert {t.name for t in teams} == {"Falcon Eye"}


@pytest.mark.asyncio
async def test_import_accepts_a_csv_export_of_the_same_tracker(db):
    """The upload widget advertises .csv as well as .xlsx — a Google Sheets
    "Download as CSV" of the same tracker tab must work identically,
    BOM and all (Excel/Sheets stamp a UTF-8 BOM on CSV exports)."""
    cohort = await _make_cohort(db)
    await db.commit()

    email = f"csvrow-{uuid.uuid4().hex[:6]}@example.com"
    csv_text = f"Team Name,Team ID,Student Name,Email,Phone,School / University\r\nFalcon Eye,T1,Amir Salem,{email},0501234567,SABIS\r\n"
    file_bytes = b"\xef\xbb\xbf" + csv_text.encode("utf-8")  # UTF-8 BOM, as Excel/Sheets stamp it

    rows = parse_registration_sheet(file_bytes, filename="eligible.csv")
    results = await import_registrations(db, cohort_id=cohort.id, rows=rows)
    await db.commit()

    assert summarize(results) == {"total": 1, "created": 1, "matched": 0, "error": 0}
    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert {t.name for t in teams} == {"Falcon Eye"}


def test_parse_rejects_a_non_spreadsheet_file_cleanly():
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as exc_info:
        parse_registration_sheet(b"this is plainly not a zip file", filename="roster.xlsx")
    assert exc_info.value.status_code == 422


@pytest.mark.asyncio
async def test_import_finds_the_header_past_the_trackers_title_and_stage_rows(db):
    """Ops's real "Eligible" export (the one that actually broke): a title
    row, a stage-description row, THEN the header on row 3 — not row 1.
    Real shape, real columns, straight from the file that failed. This
    sheet carries both Team Name and Team ID — the real, common case."""
    cohort = await _make_cohort(db)
    await db.commit()

    email = f"ed1-{uuid.uuid4().hex[:6]}@example.com"
    csv_text = (
        "Eligible Applicants — Before Qualifier,,,,,,,,,,,,,\r\n"
        "Stage 2: Confirm eligibility before the qualifier; eligible candidates are sent for NSA approval.,,,,,,,,,,,,,\r\n"
        "Application ID,Student Name,Nationality,Email,Phone,School / University,Age / Grade,Emirate,Team Name,Team ID,"
        "Eligibility Status,Eligibility Check,Eligible?,Reason / Notes,Sent to NSA?\r\n"
        f"1,ed1,Emirati,{email},,msa,12,yes,edos,1,,,,,\r\n"
    )
    file_bytes = csv_text.encode("utf-8")

    rows = parse_registration_sheet(file_bytes, filename="eligible.csv")
    results = await import_registrations(db, cohort_id=cohort.id, rows=rows)
    await db.commit()

    assert summarize(results) == {"total": 1, "created": 1, "matched": 0, "error": 0}
    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert {t.name for t in teams} == {"edos"}


@pytest.mark.asyncio
async def test_import_adds_pre_existing_user_to_team(db):
    """A student who already has a User account (e.g. from a previous LMS
    enrolment) but whose User.contact_id doesn't match the freshly-resolved
    contact should still be added to the team — not silently dropped with
    an error row (the bug: `sync_registration_lms` returns None because the
    email is already on a different User, and _import_row used to bail
    before reaching the team-join code)."""
    from app.core.security import get_password_hash

    cohort = await _make_cohort(db)

    # Create a pre-existing User with no contact_id — simulates an account
    # that was provisioned through some other path (manual admin creation,
    # an older import pipeline, etc.)
    email = f"existing-{uuid.uuid4().hex[:6]}@example.com"
    existing_user = User(
        id=uuid.uuid4(), full_name="Existing Student", email=email,
        password_hash=get_password_hash("dummy"), roles=["student"],
        status="active",
    )
    db.add(existing_user)
    await db.commit()

    file_bytes = _workbook([
        {"team_id": "T1", "team_name": "Falcon Eye", "member_name": "Existing Student", "email": email},
    ])
    rows = parse_registration_sheet(file_bytes)
    results = await import_registrations(db, cohort_id=cohort.id, rows=rows)
    await db.commit()

    # The row should succeed (not error) and the user should be on the team
    summary = summarize(results)
    assert summary["error"] == 0, f"Expected no errors, got: {[r for r in results if r.disposition == 'error']}"
    assert summary["created"] + summary["matched"] == 1

    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert len(teams) == 1
    assert teams[0].name == "Falcon Eye"
    members = await team_member_ids(db, team_id=teams[0].id)
    assert existing_user.id in members


# ── Assign existing (self-registered) students (2026-09-14) ─────────────

async def _make_self_registered_student(db, *, full_name: str = "Self Registered", email: str | None = None) -> User:
    """Mirrors what routers/auth.py::student_signup actually leaves behind:
    a Contact + a User with `contact_id` set — no Registration in any
    cohort yet. That's exactly the gap `assign_existing_students` closes."""
    from app.core.security import get_password_hash

    email = email or f"self-{uuid.uuid4().hex[:6]}@example.com"
    contact = Contact(id=uuid.uuid4(), full_name=full_name, email=email, contact_roles=["student"])
    db.add(contact)
    await db.flush()
    user = User(
        id=uuid.uuid4(), full_name=full_name, email=email, password_hash=get_password_hash("dummy"),
        roles=["student"], status="active", contact_id=contact.id, must_change_password=False,
    )
    db.add(user)
    await db.flush()
    return user


@pytest.mark.asyncio
async def test_assign_registers_a_self_registered_student_into_a_new_team(db):
    cohort = await _make_cohort(db)
    student = await _make_self_registered_student(db)
    await db.commit()

    results = await assign_existing_students(
        db, cohort_id=cohort.id, rows=[{"row_number": 1, "user_id": student.id, "team_id": "T1"}],
    )
    await db.commit()

    assert summarize_assign_results(results) == {"total": 1, "created": 1, "matched": 0, "error": 0}
    assert results[0].registration_id is not None

    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert len(teams) == 1
    assert teams[0].roster_team_id == "T1"
    assert teams[0].name == "T1"  # unnamed placeholder — student names it later
    assert student.id in await team_member_ids(db, team_id=teams[0].id)

    reg = (await db.execute(
        select(Registration).where(Registration.contact_id == student.contact_id, Registration.cohort_id == cohort.id)
    )).scalar_one()
    assert reg.id == results[0].registration_id


@pytest.mark.asyncio
async def test_assign_joins_an_existing_team_by_id_and_is_idempotent(db):
    cohort = await _make_cohort(db)
    a = await _make_self_registered_student(db, full_name="Student A")
    b = await _make_self_registered_student(db, full_name="Student B")
    await db.commit()

    first = await assign_existing_students(
        db, cohort_id=cohort.id, rows=[{"row_number": 1, "user_id": a.id, "team_id": "T1"}],
    )
    await db.commit()
    assert summarize_assign_results(first)["created"] == 1

    second = await assign_existing_students(
        db, cohort_id=cohort.id, rows=[{"row_number": 1, "user_id": b.id, "team_id": "T1"}],
    )
    await db.commit()
    assert summarize_assign_results(second) == {"total": 1, "created": 1, "matched": 0, "error": 0}

    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert len(teams) == 1
    members = set(await team_member_ids(db, team_id=teams[0].id))
    assert {a.id, b.id} <= members

    # Re-running the exact same (student, team) pair reports `matched`, not
    # a duplicate registration or a second join.
    third = await assign_existing_students(
        db, cohort_id=cohort.id, rows=[{"row_number": 1, "user_id": a.id, "team_id": "T1"}],
    )
    await db.commit()
    assert summarize_assign_results(third) == {"total": 1, "created": 0, "matched": 1, "error": 0}
    teams_again = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert len(teams_again) == 1


@pytest.mark.asyncio
async def test_assign_a_row_with_no_team_id_errors_without_failing_the_batch(db):
    cohort = await _make_cohort(db)
    a = await _make_self_registered_student(db)
    b = await _make_self_registered_student(db)
    await db.commit()

    results = await assign_existing_students(
        db, cohort_id=cohort.id,
        rows=[
            {"row_number": 1, "user_id": a.id, "team_id": ""},
            {"row_number": 2, "user_id": b.id, "team_id": "T1"},
        ],
    )
    await db.commit()

    summary = summarize_assign_results(results)
    assert summary["error"] == 1 and summary["created"] == 1
    error_row = next(r for r in results if r.disposition == "error")
    assert "Team ID" in error_row.reason


@pytest.mark.asyncio
async def test_assign_an_unknown_user_id_errors_without_failing_the_batch(db):
    cohort = await _make_cohort(db)
    real = await _make_self_registered_student(db)
    await db.commit()

    results = await assign_existing_students(
        db, cohort_id=cohort.id,
        rows=[
            {"row_number": 1, "user_id": uuid.uuid4(), "team_id": "T1"},
            {"row_number": 2, "user_id": real.id, "team_id": "T1"},
        ],
    )
    await db.commit()

    summary = summarize_assign_results(results)
    assert summary["error"] == 1 and summary["created"] == 1
    error_row = next(r for r in results if r.disposition == "error")
    assert "not found" in error_row.reason.lower()


@pytest.mark.asyncio
async def test_assign_flags_a_team_name_collision_without_failing_the_batch(db):
    """A student assigned into a fresh team_id whose fallback placeholder
    name (the team_id itself) collides case-insensitively with an existing
    real team name — same class of conflict the CSV path guards against."""
    from app.core.security import get_password_hash

    cohort = await _make_cohort(db)
    creator = User(
        id=uuid.uuid4(), full_name="Team Creator", email=f"creator-{uuid.uuid4().hex[:6]}@example.com",
        password_hash=get_password_hash("dummy"), roles=["student"], status="active",
    )
    db.add(creator)
    await db.flush()
    existing = await create_team(db, name="Falcon", created_by=creator.id, cohort_id=cohort.id)
    student = await _make_self_registered_student(db)
    await db.commit()

    results = await assign_existing_students(
        db, cohort_id=cohort.id, rows=[{"row_number": 1, "user_id": student.id, "team_id": "FALCON"}],
    )
    await db.commit()

    assert summarize_assign_results(results) == {"total": 1, "created": 0, "matched": 0, "error": 1}
    assert "already exists" in results[0].reason
    teams = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalars().all()
    assert len(teams) == 1
    assert teams[0].id == existing.id


@pytest.mark.asyncio
async def test_assign_auto_creates_contact_for_student_with_no_contact(db):
    """Per the 2026-10 rebuild in `_assign_row`, a User with `contact_id=None`
    has a Contact auto-created via `resolve_or_create_contact` so `register()`
    can attach cleanly."""
    from app.core.security import get_password_hash

    cohort = await _make_cohort(db)
    orphan = User(
        id=uuid.uuid4(), full_name="No Contact", email=f"orphan-{uuid.uuid4().hex[:6]}@example.com",
        password_hash=get_password_hash("dummy"), roles=["student"], status="active",
    )
    db.add(orphan)
    await db.commit()

    results = await assign_existing_students(
        db, cohort_id=cohort.id, rows=[{"row_number": 1, "user_id": orphan.id, "team_id": "T1"}],
        send_welcome_email=False,
    )
    await db.commit()
    assert summarize_assign_results(results) == {"total": 1, "created": 1, "matched": 0, "error": 0}
    await db.refresh(orphan)
    assert orphan.contact_id is not None


@pytest.mark.asyncio
async def test_import_loads_core_university_competition_sheet_columns_and_ignores_extra_fields(db):
    """Only the core portal fields (Student Name, Email, Phone, University ->
    Organization, Nationality, Team Name, Team ID) are loaded from the
    university competition registration sheet; extra columns (Student ID,
    University Supervisor Name/Email, Major, Year of Study, Age, University
    Campus) are safely ignored."""
    from app.models.spine.organization import Organization
    from app.services.competitions.registration_import import preview_registration_sheet

    cohort = await _make_cohort(db)
    await db.commit()

    email = f"student-{uuid.uuid4().hex[:6]}@ku.ac.ae"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Title Row"])
    sheet.append([
        "Student ID",
        "Student Name",
        "Email",
        "Phone",
        "University",
        "University Supervisor Name",
        "University Supervisor Email\n",
        "Major",
        "Year of Study",
        "Age",
        "Nationality",
        "University Campus",
        "Team Name",
        "Team ID",
    ])
    sheet.append([
        100062666.0,
        "Saeed Alshehhi",
        email,
        564459097.0,
        "Khalifa University",
        "Dr.oon",
        "boon.oi@ku.ac.ae",
        "Aerospace Engineering",
        "4th year",
        21.0,
        "United Arab Emirates",
        "Main Campus",
        "KUS06",
        1.0,
    ])
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    preview = preview_registration_sheet(file_bytes)
    assert len(preview) == 1
    p = preview[0]
    assert p.member_name == "Saeed Alshehhi"
    assert p.email == email
    assert p.phone == "564459097"
    assert p.school == "Khalifa University"
    assert p.nationality == "United Arab Emirates"
    assert p.team_name == "KUS06"
    assert p.team_id == "1"
    assert p.issues == []

    rows = parse_registration_sheet(file_bytes)
    results = await import_registrations(db, cohort_id=cohort.id, rows=rows, send_welcome_email=False)
    await db.commit()
    assert summarize(results) == {"total": 1, "created": 1, "matched": 0, "error": 0}

    contact = (await db.execute(select(Contact).where(Contact.email == email))).scalar_one()
    assert contact.full_name == "Saeed Alshehhi"
    assert contact.nationality == "United Arab Emirates"
    assert contact.primary_phone_e164 == "+971564459097"
    assert contact.organization_id is not None
    org = await db.get(Organization, contact.organization_id)
    assert org.name_latin == "Khalifa University"
    assert contact.grade is None
    assert contact.city is None
    assert contact.notes is None

    user = (await db.execute(select(User).where(User.contact_id == contact.id))).scalar_one()
    assert user.phone == "+971564459097"


@pytest.mark.asyncio
async def test_placeholder_team_is_renamed_when_subsequent_row_or_reimport_provides_team_name(db):
    """Regression test for the indentation bug where `rename_team` only ran
    inside `for other_id in other_team_ids:` — a team created with `name ==
    roster_team_id` (blank `team_name` on row 1) must be renamed when row 2 or
    a re-import supplies `team_name`, even when no member is leaving another team."""
    cohort = await _make_cohort(db)
    await db.commit()

    email_a = f"a-{uuid.uuid4().hex[:6]}@example.com"
    email_b = f"b-{uuid.uuid4().hex[:6]}@example.com"
    file_bytes = _workbook([
        {"team_id": "T-99", "team_name": None, "member_name": "Student A", "email": email_a},
        {"team_id": "T-99", "team_name": "OrbitX", "member_name": "Student B", "email": email_b},
    ])
    results = await import_registrations(
        db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes), send_welcome_email=False,
    )
    await db.commit()
    assert summarize(results) == {"total": 2, "created": 2, "matched": 0, "error": 0}

    team = (await db.execute(select(Team).where(Team.cohort_id == cohort.id))).scalar_one()
    assert team.roster_team_id == "T-99"
    assert team.name == "OrbitX"


@pytest.mark.asyncio
async def test_teammates_sharing_phone_do_not_create_unlinked_duplicate_contacts_on_reimport(db):
    """When two teammates share a phone number, `evaluate()` returns `REVIEW`
    for the second teammate. Re-importing the sheet must still match the second
    teammate by email rather than minting a third unlinked Contact."""
    cohort = await _make_cohort(db)
    await db.commit()

    email_1 = f"tm1-{uuid.uuid4().hex[:6]}@example.com"
    email_2 = f"tm2-{uuid.uuid4().hex[:6]}@example.com"
    shared_phone = "0509998877"
    file_bytes = _workbook([
        {"team_id": "T1", "team_name": "SharedPhoneTeam", "member_name": "Teammate One", "email": email_1, "phone": shared_phone},
        {"team_id": "T1", "team_name": "SharedPhoneTeam", "member_name": "Teammate Two", "email": email_2, "phone": shared_phone},
    ])
    first = await import_registrations(
        db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes), send_welcome_email=False,
    )
    await db.commit()
    assert summarize(first) == {"total": 2, "created": 2, "matched": 0, "error": 0}

    second = await import_registrations(
        db, cohort_id=cohort.id, rows=parse_registration_sheet(file_bytes), send_welcome_email=False,
    )
    await db.commit()
    assert summarize(second) == {"total": 2, "created": 0, "matched": 2, "error": 0}

    contacts_2 = (await db.execute(select(Contact).where(Contact.email == email_2))).scalars().all()
    assert len(contacts_2) == 1
    user_2 = (await db.execute(select(User).where(User.email == email_2))).scalar_one()
    reg_2 = (await db.execute(
        select(Registration).where(Registration.cohort_id == cohort.id, Registration.contact_id == contacts_2[0].id)
    )).scalar_one()
    assert user_2.contact_id == reg_2.contact_id


@pytest.mark.asyncio
async def test_preview_reports_cross_row_and_cohort_team_collisions(db):
    from app.core.security import get_password_hash
    from app.services.competitions.registration_import import preview_registration_sheet_for_cohort

    cohort = await _make_cohort(db)
    creator = User(
        id=uuid.uuid4(), full_name="Creator", email=f"creator-{uuid.uuid4().hex[:6]}@example.com",
        password_hash=get_password_hash("dummy"), roles=["student"], status="active",
    )
    db.add(creator)
    await db.flush()
    await create_team(db, name="ExistingTeam", created_by=creator.id, cohort_id=cohort.id, roster_team_id="OLD-1")
    await db.commit()

    file_bytes = _workbook([
        {"team_id": "T1", "team_name": "SameName", "member_name": "A", "email": "a@example.com"},
        {"team_id": "T2", "team_name": "SameName", "member_name": "B", "email": "b@example.com"},
        {"team_id": "T3", "team_name": "ExistingTeam", "member_name": "C", "email": "c@example.com"},
    ])
    preview = await preview_registration_sheet_for_cohort(db, cohort_id=cohort.id, file_bytes=file_bytes)
    assert preview[0].issues == []
    assert any("already exists in this sheet" in i for i in preview[1].issues)
    assert any("already exists in this cohort" in i for i in preview[2].issues)

